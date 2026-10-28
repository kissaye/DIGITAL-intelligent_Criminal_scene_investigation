from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION

OUT = '刑侦三维案例平台 安全合规与竞赛材料包.docx'

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)
def border(cell):
    tcPr=cell._tc.get_or_add_tcPr(); b=OxmlElement('w:tcBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); b.append(e)
    tcPr.append(b)
def font(run, size=10.5, bold=False, color=None):
    run.font.name='Microsoft YaHei'; run._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); run.font.size=Pt(size); run.bold=bold
    if color: run.font.color.rgb=RGBColor.from_string(color)
def add_p(doc, text='', style=None, boldlead=None):
    p=doc.add_paragraph(style=style); p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.35
    if boldlead and text.startswith(boldlead):
        font(p.add_run(boldlead),10.5,True); font(p.add_run(text[len(boldlead):]))
    else: font(p.add_run(text), 10.5)
    return p
def table(doc, headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=''; shade(c,'17365D'); border(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(p.add_run(h),9,True,'FFFFFF')
        if widths: c.width=Cm(widths[i])
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            c=cells[i]; c.text=''; border(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2==1: shade(c,'F3F7FB')
            p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(2); p.paragraph_format.space_before=Pt(2)
            font(p.add_run(val),8.8)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return t

doc=Document(); sec=doc.sections[0]; sec.top_margin=Cm(1.8); sec.bottom_margin=Cm(1.8); sec.left_margin=Cm(2.0); sec.right_margin=Cm(2.0)
styles=doc.styles
styles['Normal'].font.name='Microsoft YaHei'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); styles['Normal'].font.size=Pt(10.5)
for s,sz,col in [('Title',22,'000000'),('Heading 1',15,'000000'),('Heading 2',12,'000000')]:
    styles[s].font.name='Microsoft YaHei'; styles[s]._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); styles[s].font.size=Pt(sz); styles[s].font.color.rgb=RGBColor.from_string(col)

p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(p.add_run('刑侦三维案例平台 安全合规与竞赛材料包'),22,True)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(p.add_run('案件授权 脱敏 保密 数据生命周期 答辩与商业计划'),11,False,'404040')
doc.add_paragraph()
add_p(doc,'本材料用于支撑“数智刑勘”项目在竞赛展示、试点洽谈和内部立项中的安全合规叙事。结论是：项目的竞争力不是获取更多案件数据，而是把每一份可用数据都置于可授权、可审查、可追溯、可退出的治理闭环中。任何教学使用均不得以实战办案需要替代单独授权。')

doc.add_heading('一 参赛口径与边界',1)
table(doc,['可讲的价值','必须避免的表述','建议替代表述'],[
['一次采集服务实战留存和经批准的教学转化','“真实案件可直接进入课堂”','“仅限完成逐案授权和脱敏复核的教学副本进入受控课堂”'],
['三维数据辅助复盘、培训和可视化呈现','“三维模型天然具有证据效力”','“原始采集、保管和使用须遵循办案机关流程；平台提供可核验的数据管理能力”'],
['技术与流程共同降低泄露风险','“脱敏后绝对安全”','“通过最小化、去标识、场景处理、访问控制和复核降低再识别风险”'],
['公安与院校产教融合','“平台拥有案件数据”','“案件数据控制权和审批权始终由授权单位保留”'],
], [4.2,5.2,7.6])
add_p(doc,'底线规则：在侦查、审查起诉、审判或其他法定保密要求未解除期间，原始案件材料不进入教学库；涉及国家秘密、商业秘密、个人隐私或可能危及人员安全的内容，一律按更高等级处置。')

doc.add_heading('二 案件授权与准入机制',1)
add_p(doc,'实行“一案一档、一案一审、一案一授权”。授权不是一份泛化合作协议，而是由数据控制单位对具体案件、具体数据集、具体用途、具体期限和具体接收方作出的书面决定。平台不得自行认定案件是否可教学化。')
table(doc,['关口','责任主体','必备材料','放行条件'],[
['G0 立项','项目负责人','数据处理清单、角色清单、初步风险分类','确定仅处理必要字段与数据形态'],
['G1 原始采集','办案单位','采集任务单、设备与操作者记录、校验信息','原始数据进入实战隔离区，不自动复制'],
['G2 教学转化申请','办案单位业务部门','逐案转化申请、案件状态说明、用途和范围','确认不存在未解除的保密限制或不利影响'],
['G3 合规复核','法制 保密 数据保护负责人','脱敏方案、影响评估、复核意见','书面批准教学副本的范围、期限、接收院校'],
['G4 院校启用','院校管理员','保密承诺、课程名单、账号清单、使用规则','最小账户集开通且完成培训'],
['G5 期满退出','数据管理员','使用报告、到期清单、删除或续期申请','续期须重新审批；否则销毁教学副本并留存证明'],
], [2.4,2.7,6.0,5.9])
add_p(doc,'建议在授权文件中固定八项：数据项与版本号、案件状态、授权目的、使用主体及地点、可访问人员、有效期、禁止事项、违约处置及销毁/返还方式。涉及向院校提供的，应明确其仅为受托或获准使用方，不得二次传播、导出或用于模型训练。')

doc.add_heading('三 脱敏与教学化处理',1)
table(doc,['对象','优先处理方式','复核要点'],[
['人脸、声纹、身份证件、手机号、车牌','遮挡、替换、删除；声纹不作为教学资产保留','三维漫游、截图、导出文件中均不可还原'],
['地址、坐标、路径、时间线','模糊化、分级显示、替换为教学坐标','组合信息不会定位到真实地点或当事人'],
['当事人关系、案情文字、卷宗附件','最小化改写，使用角色代号；删除无关信息','避免通过稀有事件、职业、社区等间接识别'],
['物证编号、警号、办案单位内部标识','重新编号、映射表隔离保存或销毁','教学副本不能反向关联原案卷'],
['三维模型和全景影像','敏感区域裁切、纹理模糊、重新渲染、导出限制','评估场景本身是否仍可识别住宅或人员'],
], [3.7,6.2,7.1])
add_p(doc,'脱敏采用“双人复核”：制作人完成清单化处理，独立复核人从再识别角度检查三维场景、纹理、元数据、文件名、日志和可下载附件。复核未通过即退回，不以“已做马赛克”替代风险判断。')

doc.add_heading('四 保密与技术控制',1)
table(doc,['控制域','最低控制措施','竞赛可展示的证据'],[
['组织','数据控制方、处理方、院校管理员三级责任；保密承诺和年度培训','RACI表、培训签到、保密协议模板'],
['身份与权限','实名账号、最小权限、MFA、按课程和案例授权、离校自动回收','权限矩阵、账号开通审批单'],
['存储与传输','原始区与教学区物理或逻辑隔离；传输与静态加密；密钥分权','网络拓扑、加密和备份策略'],
['使用','仅限受控终端/网络；水印；禁止批量导出、截屏和外发，或对高敏场景禁用','水印样例、导出审批与DLP规则'],
['审计与应急','访问、检索、漫游、导出、管理员操作全留痕；异常告警；演练与处置预案','审计样例、事件演练记录、工单闭环'],
], [3.0,7.0,7.0])
add_p(doc,'发生疑似泄露时，先止损再调查：立即冻结相关账号和分享链接、保全日志和版本、评估涉及对象与影响、按授权文件和适用规定向数据控制方报告；对需通知或报告的情形，由法务/数据保护负责人统一决策与执行。')

doc.add_heading('五 数据生命周期与退出',1)
table(doc,['阶段','数据形态','核心动作','保留与销毁'],[
['采集','原始点云、影像、设备日志','校验、哈希、原始区隔离、记录采集链路','以办案机关制度为准；平台不得单方处置'],
['实战使用','原始材料与工作成果','按办案权限访问；全程审计','不得因教学需求复制'],
['转化','受控副本','授权核验、最小化、脱敏、双人复核、版本标识','未获批副本及时清除'],
['教学使用','教学案例包','课程级授权、水印、禁导出、行为审计','仅在授权期限和场所内使用'],
['归档与退出','教学副本及审计证据','到期提醒、续期复审、不可恢复删除、销毁证明','审计和授权记录按适用要求保存'],
], [2.4,3.5,7.2,3.9])
add_p(doc,'建议把“可删除性”写入产品验收指标：能够按案例、班级、用户、终端和备份副本追踪数据位置；到期后执行逻辑冻结、物理删除或密钥销毁，并输出可核验的处置记录。')

doc.add_heading('六 竞赛答辩三分钟讲稿',1)
add_p(doc,'我们解决的不是“如何把案件搬进课堂”，而是“如何让经批准的实战经验安全地进入课堂”。刑侦现场数字化采集后，原始数据始终留在办案体系；只有在案件主管单位逐案书面授权、完成用途审查和双人脱敏复核后，才生成不可反向关联原案的教学副本。')
add_p(doc,'平台以双轨架构服务两个目标：实战轨保证原始数据的采集、保全和审计；教学轨只提供受控案例、课程级授权和全程留痕。我们的商业价值不只是一台设备，而是把一次采集转化为可重复使用的合规教学资源，降低院校自建仿真场景的成本，同时不改变案件数据控制权。')
add_p(doc,'最重要的竞争壁垒是治理能力：一案一授权、一案一档、分级分域、到期退出。我们不会承诺所有案件都能教学化；恰恰相反，宁可少做，也要让每一个进入课堂的案例经得起授权、保密和审计检验。')

doc.add_heading('七 评委高频问题与回答',1)
table(doc,['问题','回答要点'],[
['真实案件能否进入教学？','只有经过办案单位逐案书面授权、案件状态与保密限制复核、脱敏及独立复核的教学副本可以进入；原始数据不进教学库。'],
['脱敏后是否就没有风险？','不能绝对化。我们采用数据最小化、场景处理、元数据清理、双人复核、限制导出和访问审计，持续管理再识别风险。'],
['平台会不会影响证据效力？','平台不替代法定取证与保管流程。原始数据的采集、校验、保管和使用由办案单位制度控制；教学副本不作为原始证据使用。'],
['商业模式是否依赖出售案件数据？','不依赖。收入来自合规部署、硬件/集成、平台订阅、课程与实训服务；案例的使用权受授权限定，不形成可自由交易的数据资产。'],
['如何防止学生传播？','实名账号、课程白名单、受控网络/终端、水印、禁导出、审计告警、承诺和违规处置共同作用；高风险案例可只在教师演示端使用。'],
], [4.4,12.6])

doc.add_heading('八 商业计划要点',1)
table(doc,['模块','竞赛表述'],[
['客户与付费方','公安刑侦及技术部门购买数字化采集与管理能力；公安、政法院校购买案例库、实训平台、课程共建和运维服务。'],
['产品包','实战采集包、数据中枢与合规模块、教学实训包、课程与师资服务。合规模块是标配，不是可选附加项。'],
['收入结构','设备/集成项目收入 + 平台订阅与运维 + 教学内容与培训服务。避免把“案件数据交易”作为收入来源。'],
['试点路径','先以2至3个授权单位和3至5所院校开展小范围、可审计试点；用合规通过率、课程使用率、续费率和零重大事件衡量。'],
['核心指标','授权完备率、脱敏复核通过率、到期处置完成率、权限回收时效、审计覆盖率、客户续费率。'],
['风险与对冲','授权不足则不上线；案件敏感度高则降级为模拟案例；技术风险以隔离、加密、审计和演练降低；法律与保密问题由控制方审批机制兜底。'],
], [3.4,13.6])

doc.add_heading('九 合规依据与使用提示',1)
add_p(doc,'本材料是竞赛和产品治理框架，不构成对具体案件、数据分类或证据资格的法律意见。试点前应由授权单位法制、保密、网安/数据管理部门及项目法务结合案件状态、数据来源、适用制度完成审查。')
for x in [
'《中华人民共和国个人信息保护法》：敏感个人信息包含生物识别、行踪轨迹等；处理敏感个人信息以及对外提供等情形应开展个人信息保护影响评估并留存记录。来源：https://www.cac.gov.cn/2021-08/20/c_1631050028355286.htm',
'《中华人民共和国数据安全法》：数据处理活动应建立全流程数据安全管理制度，并采取技术和其他必要措施保障安全。来源：https://www.cac.gov.cn/2021-06/11/c_1624994566919140.htm',
'《中华人民共和国刑事诉讼法》：对涉及国家秘密、商业秘密、个人隐私的证据应保密；技术侦查获取的材料受专门用途限制。来源：https://www.samr.gov.cn/zw/zfxxgk/fdzdgknr/bgt/art/2025/art_48b3b0c3f954448e996602a71fc6f878.html',
'《公安机关办理刑事案件程序规定》：对涉及国家秘密、商业秘密、个人隐私的证据应保密。来源：https://big5.www.gov.cn/gate/big5/www.gov.cn/zhengce/2021-12/25/content_5712867.htm'
]: add_p(doc,x)

footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(footer.add_run('刑侦三维案例平台 安全合规与竞赛材料包'),8,False,'808080')
doc.save(OUT)
print(OUT)
