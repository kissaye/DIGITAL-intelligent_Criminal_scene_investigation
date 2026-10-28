from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT='刑侦三维案例脱敏与教学实训平台 PRD V1.0.docx'
NAVY='17365D'; PALE='EAF2F8'; GRID='D9D9D9'
def font(r,s=10.5,b=False,c=None):
 r.font.name='Microsoft YaHei'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); r.font.size=Pt(s); r.font.bold=b
 if c:r.font.color.rgb=RGBColor.from_string(c)
def border(cell):
 tc=cell._tc.get_or_add_tcPr(); e=OxmlElement('w:tcBorders')
 for n in ['top','left','bottom','right','insideH','insideV']:
  x=OxmlElement('w:'+n); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),GRID); e.append(x)
 tc.append(e)
def fill(cell,c):
 e=OxmlElement('w:shd'); e.set(qn('w:fill'),c); cell._tc.get_or_add_tcPr().append(e)
def cell(c,t,h=False):
 c.text=''; p=c.paragraphs[0]; p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(2); p.paragraph_format.line_spacing=1.12
 r=p.add_run(str(t));font(r,9.1,h,'FFFFFF' if h else None);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;border(c)
 if h:fill(c,NAVY)
def tbl(d,heads,rows,widths=None):
 t=d.add_table(rows=1,cols=len(heads));t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for i,x in enumerate(heads):cell(t.rows[0].cells[i],x,True)
 trPr=t.rows[0]._tr.get_or_add_trPr(); header=OxmlElement('w:tblHeader'); header.set(qn('w:val'),'true'); trPr.append(header)
 for n,row in enumerate(rows):
  cs=t.add_row().cells
  for i,x in enumerate(row):cell(cs[i],x)
  if n%2==1:
   for c in cs:fill(c,'F7FAFC')
 if widths:
  for row in t.rows:
   for i,w in enumerate(widths):row.cells[i].width=Cm(w)
 p=d.add_paragraph();p.paragraph_format.space_after=Pt(1)
 return t
def p(d,t='',lead=None):
 x=d.add_paragraph();x.paragraph_format.line_spacing=1.42;x.paragraph_format.space_after=Pt(6)
 if lead:r=x.add_run(lead);font(r,10.5,True)
 r=x.add_run(t);font(r);return x
def h(d,t,l=1):
 x=d.add_paragraph(style='Heading '+str(l));x.paragraph_format.space_before=Pt(13 if l==1 else 8);x.paragraph_format.space_after=Pt(5)
 r=x.add_run(t);font(r,15 if l==1 else 12,True);return x
def bullets(d,items):
 for i in items:
  x=d.add_paragraph(style='List Bullet');x.paragraph_format.line_spacing=1.2;x.paragraph_format.space_after=Pt(3);r=x.add_run(i);font(r,10.2)

d=Document();s=d.sections[0];s.top_margin=Cm(2);s.bottom_margin=Cm(2);s.left_margin=Cm(2.1);s.right_margin=Cm(2.1)
d.styles['Normal'].font.name='Microsoft YaHei';d.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei');d.styles['Normal'].font.size=Pt(10.5)
x=d.add_paragraph();x.alignment=WD_ALIGN_PARAGRAPH.CENTER;x.paragraph_format.space_before=Pt(48);x.paragraph_format.space_after=Pt(14);r=x.add_run('刑侦三维案例脱敏与教学实训平台');font(r,23,True)
x=d.add_paragraph();x.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=x.add_run('产品需求文档  V1.0');font(r,13)
d.add_paragraph('\n')
tbl(d,['项目项','定义'],[
 ['产品定位','面向人工导出三维现场成果的案例脱敏、发布、教学实训和操作留痕平台。'],
 ['MVP目标','跑通可演示、可教学、可审计的案例闭环，不追求生产级公安业务系统。'],
 ['不做事项','不替代 GoSLAM；不依赖其 API；不读取 .goslam 私有采集包；不向学生开放原始案件数据。'],
 ['首期输入','LAZ/LAS/E57 点云，MP4 或全景图，案例元数据与管理员授权材料。']], [3.5,12.5])
p(d,'本版本以总控基线为唯一范围依据。所有“自动脱敏”仅作为辅助能力；任何案例发布前必须由数据管理员人工复核、确认并留痕。')

h(d,'1 业务背景与产品边界')
h(d,'1.1 要解决的问题',2)
bullets(d,['刑侦现场的三维离线成果难以在教学中安全复用：原始数据包含身份、地点、案号及其他敏感信息。','院校实训多依赖平面图片与文本，学生无法完成“看现场、找重点、测量、记录、提交、被评分”的完整训练。','竞赛原型需要证明治理闭环，而不是重复建设采集设备、公安业务系统或重型算法平台。'])
h(d,'1.2 产品边界',2)
tbl(d,['边界','平台承担','平台不承担'],[
 ['采集','接收人工导出的 LAZ/LAS/E57、MP4/全景图','设备控制、现场实时建图、读取 .goslam 文件'],
 ['脱敏','自动检测建议、人工框选、版本发布与记录','无人工审核的自动发布'],
 ['案例管理','案例包、元数据、授权范围、有效期、检索','公安案卷主系统、法定证据认定'],
 ['教学','任务、标记、测量、笔录、提交、评分与反馈','替代院校全部教务系统']], [2.2,6.5,7.3])
h(d,'1.3 MVP 成功定义',2)
p(d,'管理员可导入一个标准案例包；完成图像、点云和文本的脱敏复核；将脱敏版本发布到资源库；教师基于该案例发布实训；学生完成三维浏览、标记、测量和勘验笔录；教师评分反馈；管理员可回溯每个关键动作。')

h(d,'2 用户角色与权限')
tbl(d,['角色','核心目标','可执行动作','不可执行动作'],[
 ['数据管理员','控制案例安全与发布边界','导入、配置授权/等级/范围/有效期；复核脱敏；发布、下架；查看审计','以教师或学生身份作答'],
 ['脱敏审核员','保证脱敏结果可发布','查看检测建议；补充框选；删除/遮挡点云；提交复核','绕过管理员确认直接发布'],
 ['教师','组织与评价实训','检索已授权案例；发布任务；设物证点、重点区与评分量规；评分反馈','访问原始数据、修改脱敏版本'],
 ['学生','完成受控实训','浏览、漫游、测量、添加标记、填笔录、提交','下载原始数据、查看其他学生作答或授权材料'],
 ['平台运维','维护演示平台可用','账号、课程组织、故障处理','默认查看案例内容；需临时授权并留痕']], [2.1,3,6.5,4.4])
p(d,'权限原则：最小权限、教学副本与原始数据分离、到期自动回收、任何发布由管理员最终确认。')

h(d,'3 标准案例包与数据规范')
h(d,'3.1 目录约定',2)
tbl(d,['路径','必选','说明'],[
 ['案例包/pointcloud.laz','是','首选点云；兼容 LAS/E57 导入。'],
 ['案例包/panorama.mp4 或 images/','否','全景视频或全景图片集合。'],
 ['案例包/case.json','是','案例基础信息、分类、难度、课程标签等。'],
 ['案例包/authorization.pdf','是，仅管理员可见','授权来源、可用范围、期限等证明材料。'],
 ['案例包/desensitization_record.json','是','脱敏规则、复核人、版本、发布时间和有效期。']], [6.3,1.7,8.0])
h(d,'3.2 case json 最小字段',2)
tbl(d,['字段','说明','示例'],[
 ['case_id','内部案例编号，不使用真实案号','TC-2026-001'],['title','教学标题，不含敏感地点/人物','室内现场勘查基础训练'],['category','案件或场景类型','室内现场'],['difficulty','教学难度','初级'],['course_tags','适用课程标签','现场勘查、物证识别'],['data_level','脱敏等级','教学可用'],['valid_until','教学可用截止时间','2027-06-30']], [3.2,7.3,5.5])

h(d,'4 核心业务流程')
h(d,'4.1 案例导入与脱敏发布',2)
tbl(d,['序号','角色','动作','控制点','产物'],[
 ['1','数据管理员','上传案例包并录入授权、脱敏等级、范围和有效期','格式与必填字段校验；授权材料仅管理员可见','待处理案例'],
 ['2','系统','生成自动检测建议','图像/视频检测人脸、车牌、证件或屏幕文字；结果只作建议','待复核任务'],
 ['3','脱敏审核员','人工复核并补充处理','图像/视频模糊；点云框选删除或遮挡；文本编号化替换','脱敏候选版本'],
 ['4','数据管理员','确认发布或退回','检查脱敏记录、用途、期限和范围；不得自动发布','教学发布版本'],
 ['5','系统','转换与入库','将点云转换为网页加载格式，建立检索索引与版本记录','资源库案例'],
 ['6','系统/管理员','到期下架或续期','超过有效期自动阻断教学访问；续期需重新确认','归档或续期版本']], [1,1.8,4,6,4])
h(d,'4.2 教学实训流程',2)
tbl(d,['序号','角色','动作','结果'],[
 ['1','教师','筛选已授权案例，配置任务说明、重点区/物证点和评分量规','课程任务草稿'],
 ['2','教师','发布给指定班级并设置截止时间','学生可见的受控任务'],
 ['3','学生','在三维现场中漫游、测量、添加标记，填写勘验笔录','个人作答草稿'],
 ['4','学生','提交作答','不可改的提交版本与操作记录'],
 ['5','教师','按量规评分、批注与反馈','成绩与学习反馈'],
 ['6','管理员','查询发布、访问、提交和评分记录','可演示审计链路']], [1,2.2,6.8,6.3])
h(d,'4.3 脱敏复核闭环',2)
p(d,'自动检测建议 → 审核员查看结果 → 补充框选或调整处理 → 生成脱敏候选版本 → 数据管理员确认发布 → 教师调用教学版本 → 到期下架或重新审批。')

h(d,'5 功能需求')
tbl(d,['模块','编号','需求与验收标准','优先级'],[
 ['案例数据导入','FR-01','上传 LAZ/LAS/E57、MP4/全景图和 case.json；管理员填写授权、脱敏等级、可用范围、有效期；必填校验失败可定位。','P0'],
 ['数据脱敏','FR-02','图像/视频输出人脸、车牌、证件/屏幕文字检测建议；支持模糊与人工补框；发布前必须复核。','P1'],
 ['数据脱敏','FR-03','点云三维视图支持框选敏感区域、删除或遮挡；发布版本不含精确坐标与身份字段。','P1'],
 ['数据脱敏','FR-04','文本支持案号、人名、地址编号化替换；保存替换规则与复核记录。','P1'],
 ['三维资源库','FR-05','将获批点云转为网页可加载格式；支持案件类型、难度、课程、授权范围检索。','P0'],
 ['教育实训','FR-06','教师可创建任务、设置物证/重点区域与评分量规；学生只能使用已发布教学版本。','P0'],
 ['教育实训','FR-07','学生可漫游、测量、标记、填写勘验笔录、提交；教师可评分与反馈。','P0'],
 ['操作留痕','FR-08','记录导入、脱敏、审批、发布、访问、提交、评分等关键动作，按案例查询。','P0']], [2.5,1.4,10.0,1.5])

h(d,'6 版本规划与技术优先级')
tbl(d,['版本','目标','关键功能','完成判据'],[
 ['V0 演示原型','先证明方案能跑通','上传 LAZ/LAS、网页点云查看、点位标记、学生任务提交','一组演示数据完成“上传到提交”'],
 ['V1 核心 MVP','建成教学闭环','账号角色、案例库、教师任务、学生作答、评分、操作记录','不同角色可完成端到端教学演示'],
 ['V1.5 脱敏亮点','做出项目创新性','图像/视频人脸车牌模糊、人工审核、点云敏感区域遮挡','展示自动建议加人工确认的发布链路'],
 ['V2 完整平台','再扩展','细粒度权限、审计、案例审批、统计看板、企业/公安协作','以真实合作要求另行立项']], [2.3,2.5,7.1,5.5])
p(d,'开发顺序要求：先完成 V0 的数据可视化与学生作答，再补 V1 的教学闭环，最后以 V1.5 展示脱敏创新。任何算法实现不得阻塞 V0/V1 主链路。')

h(d,'7 非功能与风险控制')
tbl(d,['维度','MVP要求'],[
 ['数据安全','原始案件数据不进入学生端；教学端只能访问已发布副本；授权材料仅管理员可见。'],
 ['审计','核心行为具备时间、操作者、案例、版本、结果五类信息；演示可按案例回放。'],
 ['可用性','导入失败提供错误信息；大文件转换提供处理状态；失败不覆盖已发布版本。'],
 ['算法边界','检测只提供候选框和置信提示；没有人工确认，不允许发布。'],
 ['演示边界','可用模拟或合法授权演示数据；不宣称生产级公安系统能力或证据效力。']], [3,12])

h(d,'8 MVP 验收清单')
bullets(d,['管理员能导入符合目录规范的案例包，并录入授权、脱敏等级、范围和有效期。','可在网页查看点云并创建、保存重点区域/物证标记；学生能完成任务并提交笔录。','图像/视频脱敏提供自动检测建议及人工补框；点云可框选遮挡或删除敏感区；文本可编号化。','未经管理员确认的脱敏候选版本不能进入资源库和学生端。','教师可发布任务、配置量规并完成评分反馈；学生只看自己课程内的发布版本。','演示时可展示从导入、脱敏审核、发布、学生提交到教师评分的完整操作记录。'])

h(d,'9 待决策项')
tbl(d,['事项','建议','影响'],[
 ['网页点云技术栈','V0 优先选择成熟开源网页点云查看器，不自研渲染引擎','决定 V0 演示速度与 LAZ 支持方式'],
 ['点云转换格式','以浏览器可加载、可做区域遮挡的中间格式为准，保留原始 LAZ/LAS/E57','决定资源库处理队列'],
 ['脱敏模型','优先接入成熟检测能力或离线推理；无模型时保留人工框选链路','不应阻塞 V0/V1'],
 ['演示案例','使用模拟或已获教学授权数据，避免真实案件敏感信息','决定合规与展示材料']], [3.3,8.2,3.5])

for sec in d.sections:
 q=sec.footer.paragraphs[0];q.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=q.add_run('刑侦三维案例脱敏与教学实训平台 PRD V1.0');font(r,8,False,'666666')
d.save(OUT);print(OUT)
